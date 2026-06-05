#!/usr/bin/env python3
"""Raw data audit tool for UniCure and other tabular raw-data packages."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import numpy as np
import pandas as pd
from scipy import stats

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover
    Document = None
    WD_ALIGN_PARAGRAPH = None
    Inches = None
    Pt = None
    RGBColor = None

try:
    import openpyxl  # noqa: F401
except ImportError:  # pragma: no cover
    openpyxl = None

try:
    import pyreadr  # type: ignore
except ImportError:  # pragma: no cover
    pyreadr = None

try:
    import anndata  # type: ignore
except ImportError:  # pragma: no cover
    anndata = None

try:
    import h5py  # type: ignore
except ImportError:  # pragma: no cover
    h5py = None


SUPPORTED_FORMATS = {
    "csv", "tsv", "txt", "xlsx", "xls", "rdata", "rds", "rda", "parquet", "h5ad", "h5", "hdf5"
}

ROLE_PATTERNS = {
    "identifier": re.compile(r"(^|[_-])(id|sample|patient|cell|gene|drug|compound|barcode|name|type|group|label)([_-]|$)|(^|[_-])(mouse|animal)([_-]?(id|name|group|type)|$)", re.I),
    "coordinate": re.compile(r"(^|_)x$|(^|_)y$|tsne|t-sne|umap|\bpc\b|pca|dim|coordinate|coord", re.I),
    "time_dose_rank": re.compile(r"time|day|week|hour|dose|concentration|conc|epoch|step|iteration|rank|order|index|^unnamed(?::\s*\d+)?$|^\d+$", re.I),
    "statistical": re.compile(r"pvalue|p_value|p-value|p[._-]?val|p[._-]?adj|padj|fdr|qvalue|q_value|q[._-]?val|logfc|auc|pearson|spearman|correlation|corr|r2|loss|score|metric", re.I),
    "count": re.compile(r"count|reads|umi|n_cells|num|number|total", re.I),
}

RAW_PVALUE_NAME_RE = re.compile(r"(^|[_\-.\s])(p|pvalue|pval|p_value|p-value|p\.value|p_val|padj|p_adj|p-adj|p\.adj|fdr|qvalue|q_value|q-value|qval|q_val)([_\-.\s]|$)", re.I)
TRANSFORMED_PVALUE_NAME_RE = re.compile(
    r"(-|−|minus|neg|negative)?\s*log(?:10|2|e)?\s*[_\-.\s]*p(?:value|val|adj)?|"
    r"(p(?:value|val|adj)?|padj|qvalue|qval|fdr)\s*[_\-.\s]*(-|−|minus|neg|negative)?\s*log(?:10|2|e)?",
    re.I,
)

NA_STRINGS = {"", "na", "nan", "none", "null", "missing", "n/a", "NA", "NaN", "NULL"}


@dataclass
class AuditConfig:
    raw_dir: Path
    output_docx: Path
    output_json: Path | None
    formats: set[str]
    max_file_mb: float
    max_rows: int
    max_cols: int
    sample_rows: int
    min_numeric_n: int
    alpha: float
    yellow_threshold: float
    red_threshold: float
    include_hidden: bool
    verbose: bool
    run_simulations: bool
    audit_simulations_and_raw: bool
    simulation_output_dir: Path
    no_docx: bool
    no_json: bool


@dataclass
class DependencyStatus:
    pandas: bool = True
    numpy: bool = True
    scipy: bool = True
    python_docx: bool = Document is not None
    openpyxl: bool = openpyxl is not None
    pyreadr: bool = pyreadr is not None
    anndata: bool = anndata is not None
    h5py: bool = h5py is not None


@dataclass
class LoadedTable:
    file_path: Path
    file_format: str
    table_name: str
    dataframe: pd.DataFrame | None
    status: str
    notes: list[str] = field(default_factory=list)
    n_rows: int | None = None
    n_cols: int | None = None


@dataclass
class Finding:
    check_name: str
    family: str
    severity: str
    confidence: str
    score: float
    message: str
    details: dict[str, Any] = field(default_factory=dict)
    false_positive_notes: list[str] = field(default_factory=list)
    column: str | None = None


@dataclass
class TableAuditResult:
    table_name: str
    status: str
    n_rows: int | None
    n_cols: int | None
    numeric_cols: int
    missing_fraction: float | None
    duplicate_fraction: float | None
    roles: dict[str, list[str]]
    findings: list[Finding]
    notes: list[str]


@dataclass
class FileAuditResult:
    path: str
    format: str
    size_mb: float
    status: str
    category: str
    score: float
    tables: list[TableAuditResult]
    notes: list[str]
    content_hash: str | None = None


@dataclass
class ProgressEvent:
    phase: str
    current: int
    total: int
    path: str | None = None
    message: str = ""


ProgressCallback = Callable[[ProgressEvent], None]


def emit_progress(
    callback: ProgressCallback | None,
    phase: str,
    current: int,
    total: int,
    path: str | None = None,
    message: str = "",
) -> None:
    if callback is not None:
        callback(ProgressEvent(phase, current, total, path, message))


@dataclass
class SimulationExpectedResult:
    name: str
    path: str
    expected_category: str
    observed_category: str
    passed: bool
    notes: str


def parse_args() -> AuditConfig:
    parser = argparse.ArgumentParser(
        description="Audit raw-data files for statistical and structural integrity indicators."
    )
    default_out = Path("raw_data_check_soft") / "raw_data_audit_report.docx"
    default_json = Path("raw_data_check_soft") / "raw_data_audit_report.json"
    parser.add_argument("--raw-dir", type=Path, default=Path("raw_data"))
    parser.add_argument("--output-docx", type=Path, default=default_out)
    parser.add_argument("--output-json", type=Path, default=default_json)
    parser.add_argument("--formats", default=",".join(sorted(SUPPORTED_FORMATS)))
    parser.add_argument("--max-file-mb", type=float, default=250.0)
    parser.add_argument("--max-rows", type=int, default=200000)
    parser.add_argument("--max-cols", type=int, default=5000)
    parser.add_argument("--sample-rows", type=int, default=100000)
    parser.add_argument("--min-numeric-n", type=int, default=50)
    parser.add_argument("--alpha", type=float, default=0.01)
    parser.add_argument("--yellow-threshold", type=float, default=25.0)
    parser.add_argument("--red-threshold", type=float, default=90.0)
    parser.add_argument("--include-hidden", action="store_true")
    parser.add_argument("--verbose", action="store_true")
    parser.add_argument("--run-simulations", action="store_true")
    parser.add_argument("--audit-simulations-and-raw", action="store_true")
    parser.add_argument("--simulation-output-dir", type=Path, default=Path("raw_data_check_soft") / "simulation_data")
    parser.add_argument("--no-docx", action="store_true")
    parser.add_argument("--no-json", action="store_true")
    args = parser.parse_args()
    formats = {x.strip().lower().lstrip(".") for x in args.formats.split(",") if x.strip()}
    return AuditConfig(
        raw_dir=args.raw_dir,
        output_docx=args.output_docx,
        output_json=None if args.no_json else args.output_json,
        formats=formats,
        max_file_mb=args.max_file_mb,
        max_rows=args.max_rows,
        max_cols=args.max_cols,
        sample_rows=args.sample_rows,
        min_numeric_n=args.min_numeric_n,
        alpha=args.alpha,
        yellow_threshold=args.yellow_threshold,
        red_threshold=args.red_threshold,
        include_hidden=args.include_hidden,
        verbose=args.verbose,
        run_simulations=args.run_simulations,
        audit_simulations_and_raw=args.audit_simulations_and_raw,
        simulation_output_dir=args.simulation_output_dir,
        no_docx=args.no_docx,
        no_json=args.no_json,
    )


def file_format(path: Path) -> str:
    suffix = path.suffix.lower().lstrip(".")
    if suffix in {"rdata", "rda"}:
        return "rdata"
    return suffix


def is_hidden(path: Path) -> bool:
    return any(part.startswith(".") for part in path.parts)


def discover_files(root: Path, config: AuditConfig) -> list[Path]:
    if not root.exists():
        return []
    files: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if not config.include_hidden and is_hidden(path.relative_to(root)):
            continue
        fmt = file_format(path)
        if fmt in config.formats:
            files.append(path)
    return sorted(files)


def content_hash(path: Path, limit_mb: float = 100.0) -> str | None:
    if path.stat().st_size > limit_mb * 1024 * 1024:
        return None
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_csv_like(path: Path, sep: str | None, config: AuditConfig) -> LoadedTable:
    encodings = ["utf-8", "utf-8-sig", "latin-1"]
    notes: list[str] = []
    if sep is None:
        sep = sniff_delimiter(path) or ","
        notes.append(f"Detected delimiter: {repr(sep)}")
    for enc in encodings:
        try:
            df = pd.read_csv(path, sep=sep, encoding=enc, nrows=config.max_rows)
            df = cap_columns(df, config)
            return LoadedTable(path, file_format(path), "main", df, "loaded", notes, len(df), len(df.columns))
        except Exception as exc:
            notes.append(f"read_csv failed with encoding {enc}: {exc}")
    return LoadedTable(path, file_format(path), "main", None, "failed", notes)


def sniff_delimiter(path: Path) -> str | None:
    try:
        sample = path.read_text(encoding="utf-8", errors="ignore")[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        return dialect.delimiter
    except Exception:
        return None


def cap_columns(df: pd.DataFrame, config: AuditConfig) -> pd.DataFrame:
    if len(df.columns) > config.max_cols:
        return df.iloc[:, : config.max_cols].copy()
    return df


def sample_rows(df: pd.DataFrame, config: AuditConfig) -> pd.DataFrame:
    if len(df) > config.sample_rows:
        return df.sample(n=config.sample_rows, random_state=7).sort_index()
    return df


def trim_empty_edges(df: pd.DataFrame) -> pd.DataFrame:
    cleaned = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    return cleaned.reset_index(drop=True)


def contiguous_true_runs(mask: list[bool]) -> list[tuple[int, int]]:
    runs: list[tuple[int, int]] = []
    start: int | None = None
    for index, value in enumerate(mask):
        if value and start is None:
            start = index
        elif not value and start is not None:
            runs.append((start, index))
            start = None
    if start is not None:
        runs.append((start, len(mask)))
    return runs


def excel_column_group_starts(row_block: pd.DataFrame) -> list[int]:
    starts: list[int] = []
    for _, row in row_block.head(4).iterrows():
        for idx, value in enumerate(row.tolist()):
            if isinstance(value, str) and re.search(r"concentrat|dose", value, re.I):
                starts.append(idx)
    starts = sorted(set(starts))
    if len(starts) >= 2:
        return starts
    return []


def split_excel_layout_blocks(df: pd.DataFrame) -> list[pd.DataFrame]:
    cleaned = trim_empty_edges(df)
    if cleaned.empty:
        return []

    row_mask = cleaned.notna().any(axis=1).tolist()
    row_runs = contiguous_true_runs(row_mask)
    blocks: list[pd.DataFrame] = []

    for row_start, row_end in row_runs:
        row_block = cleaned.iloc[row_start:row_end, :]
        col_mask = row_block.notna().any(axis=0).tolist()
        for col_start, col_end in contiguous_true_runs(col_mask):
            block = trim_empty_edges(row_block.iloc[:, col_start:col_end])
            if block.shape[0] < 2 or block.shape[1] < 2:
                continue
            starts = excel_column_group_starts(block)
            if starts:
                starts_with_end = starts + [block.shape[1]]
                for idx, start_col in enumerate(starts):
                    end_col = starts_with_end[idx + 1]
                    sub = trim_empty_edges(block.iloc[:, start_col:end_col])
                    if sub.shape[0] >= 2 and sub.shape[1] >= 2:
                        blocks.append(sub)
            else:
                blocks.append(block)

    return blocks or [cleaned]


def normalize_excel_block(block: pd.DataFrame) -> pd.DataFrame:
    block = trim_empty_edges(block)
    if block.empty:
        return block
    header_scores = []
    for idx in range(min(3, len(block))):
        row = block.iloc[idx]
        text_count = sum(isinstance(value, str) and bool(value.strip()) for value in row.dropna())
        non_null = row.notna().sum()
        header_scores.append((text_count, non_null, idx))
    header_idx = max(header_scores, key=lambda x: (x[0], x[1]))[2] if header_scores else 0
    if header_idx > 0 or any(isinstance(value, str) for value in block.iloc[header_idx].dropna()):
        columns = []
        seen: Counter[str] = Counter()
        for col_idx, value in enumerate(block.iloc[header_idx].tolist()):
            base_name = str(value).strip() if pd.notna(value) and str(value).strip() else f"column_{col_idx + 1}"
            seen[base_name] += 1
            name = base_name if seen[base_name] == 1 else f"{base_name}_{seen[base_name]}"
            columns.append(name)
        data = block.iloc[header_idx + 1:].reset_index(drop=True)
        data.columns = columns
        return trim_empty_edges(data)
    block = block.reset_index(drop=True)
    block.columns = [f"column_{i + 1}" for i in range(block.shape[1])]
    return block


def load_excel(path: Path, config: AuditConfig) -> list[LoadedTable]:
    if openpyxl is None and path.suffix.lower() == ".xlsx":
        return [LoadedTable(path, file_format(path), "workbook", None, "metadata_only", ["Install openpyxl to audit .xlsx files."])]
    try:
        sheets = pd.read_excel(path, sheet_name=None, nrows=config.max_rows, header=None)
        tables: list[LoadedTable] = []
        for sheet_name, sheet_df in sheets.items():
            blocks = split_excel_layout_blocks(sheet_df)
            if len(blocks) <= 1:
                df = normalize_excel_block(blocks[0] if blocks else sheet_df)
                df = cap_columns(df, config)
                tables.append(LoadedTable(path, file_format(path), str(sheet_name), df, "loaded", ["Excel sheet was trimmed before audit."], len(df), len(df.columns)))
            else:
                for block_index, block in enumerate(blocks, start=1):
                    df = normalize_excel_block(block)
                    if df.empty:
                        continue
                    df = cap_columns(df, config)
                    tables.append(LoadedTable(path, file_format(path), f"{sheet_name} / block {block_index}", df, "loaded", ["Excel layout sheet was split into contiguous data blocks before audit."], len(df), len(df.columns)))
        return tables or [LoadedTable(path, file_format(path), "workbook", None, "metadata_only", ["No non-empty Excel data blocks found."])]
    except Exception as exc:
        return [LoadedTable(path, file_format(path), "workbook", None, "failed", [str(exc)])]


def load_rdata(path: Path, config: AuditConfig) -> list[LoadedTable]:
    if pyreadr is None:
        return [LoadedTable(path, "rdata", "R objects", None, "metadata_only", ["Install pyreadr to audit .Rdata/.rds objects."])]
    try:
        result = pyreadr.read_r(str(path))
        tables: list[LoadedTable] = []
        for name, obj in result.items():
            if isinstance(obj, pd.DataFrame):
                df = cap_columns(obj.head(config.max_rows), config)
                tables.append(LoadedTable(path, "rdata", str(name), df, "loaded", [], len(df), len(df.columns)))
            else:
                tables.append(LoadedTable(path, "rdata", str(name), None, "metadata_only", [f"Object type {type(obj).__name__} was not table-like."]))
        return tables or [LoadedTable(path, "rdata", "R objects", None, "metadata_only", ["No readable R objects found."])]
    except Exception as exc:
        return [LoadedTable(path, "rdata", "R objects", None, "failed", [str(exc)])]


def load_parquet(path: Path, config: AuditConfig) -> list[LoadedTable]:
    try:
        df = pd.read_parquet(path)
        df = df.head(config.max_rows)
        df = cap_columns(df, config)
        return [LoadedTable(path, "parquet", "main", df, "loaded", [], len(df), len(df.columns))]
    except Exception as exc:
        return [LoadedTable(path, "parquet", "main", None, "failed", [str(exc)])]


def load_h5ad(path: Path, config: AuditConfig) -> list[LoadedTable]:
    if anndata is None:
        return [LoadedTable(path, "h5ad", "AnnData", None, "metadata_only", ["Install anndata or scanpy to audit .h5ad files."])]
    try:
        ad = anndata.read_h5ad(path, backed="r")
        tables = [
            LoadedTable(path, "h5ad", "obs", cap_columns(ad.obs.head(config.max_rows).reset_index(), config), "loaded", [], ad.n_obs, len(ad.obs.columns)),
            LoadedTable(path, "h5ad", "var", cap_columns(ad.var.head(config.max_rows).reset_index(), config), "loaded", [], ad.n_vars, len(ad.var.columns)),
        ]
        return tables
    except Exception as exc:
        return [LoadedTable(path, "h5ad", "AnnData", None, "failed", [str(exc)])]


def load_h5(path: Path, config: AuditConfig) -> list[LoadedTable]:
    if h5py is None:
        return [LoadedTable(path, "h5", "HDF5", None, "metadata_only", ["Install h5py to audit .h5/.hdf5 files."])]
    tables: list[LoadedTable] = []
    try:
        with h5py.File(path, "r") as handle:
            def visitor(name: str, obj: Any) -> None:
                if hasattr(obj, "shape") and obj.shape and len(obj.shape) <= 2 and np.prod(obj.shape) <= config.max_rows * min(config.max_cols, 1000):
                    try:
                        arr = obj[()]
                        if np.issubdtype(arr.dtype, np.number):
                            df = pd.DataFrame(arr)
                            df = cap_columns(df.head(config.max_rows), config)
                            tables.append(LoadedTable(path, "h5", name, df, "loaded", [], len(df), len(df.columns)))
                    except Exception:
                        pass
            handle.visititems(visitor)
        return tables or [LoadedTable(path, "h5", "HDF5", None, "metadata_only", ["No small numeric 1D/2D datasets found."])]
    except Exception as exc:
        return [LoadedTable(path, "h5", "HDF5", None, "failed", [str(exc)])]


def load_file(path: Path, config: AuditConfig) -> list[LoadedTable]:
    fmt = file_format(path)
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > config.max_file_mb:
        return [LoadedTable(path, fmt, "file", None, "metadata_only", [f"File exceeds --max-file-mb ({size_mb:.1f} MB)."])]
    if fmt == "csv":
        return [read_csv_like(path, ",", config)]
    if fmt == "tsv":
        return [read_csv_like(path, "\t", config)]
    if fmt == "txt":
        return [read_csv_like(path, None, config)]
    if fmt in {"xlsx", "xls"}:
        return load_excel(path, config)
    if fmt in {"rdata", "rds", "rda"}:
        return load_rdata(path, config)
    if fmt == "parquet":
        return load_parquet(path, config)
    if fmt == "h5ad":
        return load_h5ad(path, config)
    if fmt in {"h5", "hdf5"}:
        return load_h5(path, config)
    return [LoadedTable(path, fmt, "file", None, "metadata_only", ["Unsupported format."])]


def normalize_for_duplicate(df: pd.DataFrame) -> pd.DataFrame:
    return df.astype(str).apply(lambda col: col.str.strip().str.lower()).replace({x.lower(): "" for x in NA_STRINGS})


def infer_roles(df: pd.DataFrame) -> dict[str, list[str]]:
    roles: dict[str, list[str]] = {}
    for col in df.columns:
        name = str(col)
        detected = [role for role, pat in ROLE_PATTERNS.items() if pat.search(name)]
        series = df[col]
        if pd.api.types.is_numeric_dtype(series):
            unique_ratio = series.nunique(dropna=True) / max(series.notna().sum(), 1)
            if unique_ratio < 0.05 and "identifier" not in detected:
                detected.append("low_cardinality")
        roles[name] = detected or ["measurement_like"]
    return roles


def numeric_columns(df: pd.DataFrame) -> list[str]:
    cols: list[str] = []
    for col in df.columns:
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.notna().sum() > 0:
            cols.append(str(col))
    return cols


def get_numeric_series(df: pd.DataFrame, col: str) -> pd.Series:
    return pd.to_numeric(df[col], errors="coerce").dropna()


def role_has(roles: dict[str, list[str]], col: str, role: str) -> bool:
    return role in roles.get(str(col), [])


def is_contextual_column(roles: dict[str, list[str]], col: str) -> bool:
    contextual = {"identifier", "coordinate", "time_dose_rank", "statistical", "count", "low_cardinality"}
    return bool(contextual.intersection(roles.get(str(col), [])))


def add_finding(findings: list[Finding], check_name: str, family: str, severity: str, confidence: str, score: float, message: str, details: dict[str, Any] | None = None, false_positive_notes: list[str] | None = None, column: str | None = None) -> None:
    findings.append(Finding(check_name, family, severity, confidence, score, message, details or {}, false_positive_notes or [], column))


def check_duplicates(df: pd.DataFrame, findings: list[Finding]) -> None:
    if df.empty:
        return
    dup_frac = float(df.duplicated().mean())
    if dup_frac >= 0.20:
        add_finding(findings, "Exact duplicate rows", "duplicates", "red", "high", 45, f"Exact duplicate rows account for {dup_frac:.1%} of the table.", {"duplicate_fraction": dup_frac})
    elif dup_frac >= 0.05:
        add_finding(findings, "Exact duplicate rows", "duplicates", "yellow", "medium", 18, f"Exact duplicate rows account for {dup_frac:.1%} of the table.", {"duplicate_fraction": dup_frac})
    norm = normalize_for_duplicate(df)
    norm_dup = float(norm.duplicated().mean())
    if norm_dup > dup_frac and norm_dup >= 0.10:
        add_finding(findings, "Normalized duplicate rows", "duplicates", "yellow", "medium", 20, f"Rows duplicated after whitespace/NA normalization account for {norm_dup:.1%}.", {"normalized_duplicate_fraction": norm_dup})
    duplicate_col_names = [k for k, v in Counter(map(str, df.columns)).items() if v > 1]
    if duplicate_col_names:
        add_finding(findings, "Duplicate column names", "duplicates", "yellow", "high", 15, "Duplicate column names were found.", {"columns": duplicate_col_names[:20]})
    if len(df) >= 20:
        run_lengths = []
        prev_hash = None
        current = 0
        for _, row in norm.iterrows():
            row_hash = hashlib.md5("|".join(map(str, row.values)).encode("utf-8", errors="ignore")).hexdigest()
            if row_hash == prev_hash:
                current += 1
            else:
                if current > 1:
                    run_lengths.append(current)
                current = 1
                prev_hash = row_hash
        if current > 1:
            run_lengths.append(current)
        if run_lengths and max(run_lengths) >= 5:
            add_finding(findings, "Contiguous repeated rows", "duplicates", "yellow", "medium", 20, f"A contiguous repeated-row run of length {max(run_lengths)} was detected.", {"max_run_length": max(run_lengths)})


def visible_terminal_digits(series: pd.Series) -> list[int]:
    digits: list[int] = []
    for value in series.dropna().astype(str):
        value = value.strip()
        if not value or value.lower() in {"nan", "none"}:
            continue
        match = re.search(r"(\d)(?:\.0+)?$", value)
        if match:
            digits.append(int(match.group(1)))
    return digits


def check_terminal_digits(df: pd.DataFrame, roles: dict[str, list[str]], findings: list[Finding], config: AuditConfig) -> None:
    pvals: list[tuple[str, float, float, float, int]] = []
    for col in numeric_columns(df):
        if is_contextual_column(roles, col):
            continue
        nums = get_numeric_series(df, col)
        if nums.size < config.min_numeric_n:
            continue
        digits = visible_terminal_digits(df[col])
        if len(digits) < config.min_numeric_n:
            scaled = np.round(np.abs(nums.to_numpy()) * 1000).astype(np.int64)
            digits = [int(x % 10) for x in scaled]
        if len(digits) < config.min_numeric_n:
            continue
        counts = np.bincount(digits, minlength=10)
        expected = np.ones(10) * (len(digits) / 10)
        chi, p = stats.chisquare(counts, expected)
        zero_five_share = (counts[0] + counts[5]) / max(counts.sum(), 1)
        entropy = digit_entropy(counts)
        pvals.append((col, float(p), float(zero_five_share), float(entropy), int(len(digits))))
    adjusted = bh_adjust([x[1] for x in pvals])
    for (col, p, zero_five_share, entropy, n), q in zip(pvals, adjusted):
        if q <= config.alpha and zero_five_share >= 0.40:
            add_finding(findings, "Terminal digit 0/5 heaping", "digit_patterns", "red" if zero_five_share >= 0.55 else "yellow", "medium", 35 if zero_five_share >= 0.55 else 20, f"Column {col} has {zero_five_share:.1%} terminal 0/5 digits.", {"column": col, "n": n, "p_value": p, "q_value": q, "zero_five_share": zero_five_share}, ["Rounded measurement protocols can create terminal-digit heaping."], col)
        elif q <= config.alpha and entropy < 0.65:
            add_finding(findings, "Low terminal-digit entropy", "digit_patterns", "yellow", "medium", 15, f"Column {col} has unusually low terminal-digit entropy.", {"column": col, "n": n, "p_value": p, "q_value": q, "entropy": entropy}, ["Small or rounded datasets can reduce digit entropy."], col)


def digit_entropy(counts: np.ndarray) -> float:
    total = counts.sum()
    if total == 0:
        return 0.0
    probs = counts[counts > 0] / total
    return float(-(probs * np.log(probs)).sum() / math.log(10))


def bh_adjust(pvalues: list[float]) -> list[float]:
    n = len(pvalues)
    if n == 0:
        return []
    order = np.argsort(pvalues)
    adjusted = np.empty(n, dtype=float)
    prev = 1.0
    for rank, idx in enumerate(order[::-1], start=1):
        original_rank = n - rank + 1
        q = min(prev, pvalues[idx] * n / original_rank)
        adjusted[idx] = q
        prev = q
    return adjusted.tolist()


def check_heaping(df: pd.DataFrame, roles: dict[str, list[str]], findings: list[Finding], config: AuditConfig) -> None:
    for col in numeric_columns(df):
        nums = get_numeric_series(df, col)
        if nums.size < config.min_numeric_n:
            continue
        unique_ratio = nums.nunique() / nums.size
        mode_share = nums.value_counts(normalize=True).iloc[0] if nums.nunique() else 1.0
        roles_for_col = set(roles.get(str(col), []))
        contextual = bool(({"identifier", "coordinate", "time_dose_rank", "statistical", "count"}).intersection(roles_for_col))
        low_cardinality_measurement = "low_cardinality" in roles_for_col and not contextual
        if (not contextual or low_cardinality_measurement) and unique_ratio < 0.10 and mode_share >= 0.15:
            severity = "yellow" if low_cardinality_measurement else "yellow"
            score = 22 if low_cardinality_measurement else 18
            add_finding(findings, "Low unique-value ratio", "heaping_low_randomness", severity, "medium", score, f"Column {col} has low unique-value ratio ({unique_ratio:.2f}) and mode share {mode_share:.1%}.", {"column": col, "unique_ratio": float(unique_ratio), "mode_share": float(mode_share)}, ["Count data, scores, and rounded measurements may naturally repeat values."], col)
        if nums.size >= 100:
            rounded_one = np.isclose(nums, np.round(nums, 1)).mean()
            integer_like = np.isclose(nums, np.round(nums)).mean()
            if not contextual and (rounded_one >= 0.95 or integer_like >= 0.80) and nums.nunique() > 10:
                add_finding(findings, "Strong rounding pattern", "heaping_low_randomness", "yellow", "low", 12, f"Column {col} is dominated by rounded values.", {"rounded_to_one_decimal_share": float(rounded_one), "integer_like_share": float(integer_like)}, ["Some instruments or published summaries intentionally round values."], col)


def check_coordinate_regularities(df: pd.DataFrame, roles: dict[str, list[str]], findings: list[Finding], config: AuditConfig) -> None:
    num_cols = [col for col in numeric_columns(df) if get_numeric_series(df, col).size >= config.min_numeric_n]
    if not num_cols or len(df) < config.min_numeric_n:
        return
    suspicious_cols = []
    row_idx = np.arange(len(df), dtype=float)
    for col in num_cols:
        nums = pd.to_numeric(df[col], errors="coerce")
        valid = nums.notna().to_numpy()
        if valid.sum() < config.min_numeric_n:
            continue
        y = nums.to_numpy(dtype=float)[valid]
        x = row_idx[valid]
        if np.std(y) == 0:
            continue
        corr = abs(np.corrcoef(x, y)[0, 1]) if len(y) > 1 else 0.0
        diffs = np.diff(y)
        repeated_delta_share = repeated_delta_fraction(diffs)
        contextual = is_contextual_column(roles, col)
        if corr >= 0.98 and not contextual:
            suspicious_cols.append((col, corr, repeated_delta_share))
        elif repeated_delta_share >= 0.90 and len(diffs) >= config.min_numeric_n and not contextual:
            suspicious_cols.append((col, corr, repeated_delta_share))
    if len(suspicious_cols) >= 3:
        add_finding(findings, "Coordinate-related column trends", "coordinate_regularities", "red", "medium", 45, f"{len(suspicious_cols)} measurement-like columns are strongly tied to row order or repeated deltas.", {"columns": suspicious_cols[:20]}, ["Sorted tables, time courses, dose responses, and ranks can legitimately show row-order structure."])
    elif len(suspicious_cols) >= 1:
        add_finding(findings, "Coordinate-related column trend", "coordinate_regularities", "yellow", "low", 15, f"{len(suspicious_cols)} measurement-like column shows strong row-order or repeated-delta structure.", {"columns": suspicious_cols[:10]}, ["Single monotonic columns can be legitimate, especially time, dose, rank, or sorted data."])
    matrix_cols = [col for col in num_cols if not is_contextual_column(roles, col)]
    if len(matrix_cols) >= 4 and len(df) >= 20:
        block = pd.DataFrame({col: pd.to_numeric(df[col], errors="coerce") for col in matrix_cols}).dropna(axis=0, how="any")
        if 20 <= len(block) <= config.sample_rows:
            r2 = coordinate_plane_r2(block.to_numpy(dtype=float))
            if r2 >= 0.95:
                add_finding(findings, "2D table-coordinate gradient", "coordinate_regularities", "red", "medium", 50, "A numeric block is strongly explained by row and column indices.", {"r2": float(r2), "n_rows": int(block.shape[0]), "n_cols": int(block.shape[1])}, ["Some designed matrices may intentionally encode gradients or ordered grids."])


def repeated_delta_fraction(diffs: np.ndarray) -> float:
    if len(diffs) == 0:
        return 0.0
    rounded = np.round(diffs, 8)
    counts = Counter(rounded)
    return max(counts.values()) / len(diffs)


def coordinate_plane_r2(arr: np.ndarray) -> float:
    rows, cols = arr.shape
    rr, cc = np.indices((rows, cols))
    X = np.column_stack([rr.ravel(), cc.ravel(), np.ones(rows * cols)])
    y = arr.ravel()
    if np.std(y) == 0:
        return 0.0
    coef, *_ = np.linalg.lstsq(X, y, rcond=None)
    pred = X @ coef
    ss_res = np.sum((y - pred) ** 2)
    ss_tot = np.sum((y - y.mean()) ** 2)
    return float(1 - ss_res / ss_tot) if ss_tot > 0 else 0.0


def check_benford(df: pd.DataFrame, roles: dict[str, list[str]], findings: list[Finding], config: AuditConfig) -> None:
    expected = np.array([math.log10(1 + 1 / d) for d in range(1, 10)])
    for col in numeric_columns(df):
        if is_contextual_column(roles, col):
            continue
        nums = get_numeric_series(df, col)
        nums = nums[nums > 0].abs()
        if nums.size < max(config.min_numeric_n, 100):
            continue
        if nums.max() / max(nums.min(), 1e-12) < 100:
            continue
        first_digits = []
        for value in nums:
            s = f"{value:.12g}".replace(".", "").lstrip("0")
            if s and s[0].isdigit() and s[0] != "0":
                first_digits.append(int(s[0]))
        if len(first_digits) < max(config.min_numeric_n, 100):
            continue
        counts = np.bincount(first_digits, minlength=10)[1:]
        chi, p = stats.chisquare(counts, expected * counts.sum())
        if p <= config.alpha:
            add_finding(findings, "Benford first-digit deviation", "benford", "yellow", "low", 12, f"Column {col} deviates from Benford first-digit expectation.", {"column": col, "p_value": float(p), "n": len(first_digits)}, ["Benford tests are only appropriate for certain unbounded, multi-scale measurements and are weak evidence alone."], col)


def is_raw_pvalue_column(name: str) -> bool:
    normalized = str(name).strip().lower()
    if TRANSFORMED_PVALUE_NAME_RE.search(normalized):
        return False
    return bool(RAW_PVALUE_NAME_RE.search(normalized))


def check_scientific_bounds(df: pd.DataFrame, roles: dict[str, list[str]], findings: list[Finding]) -> None:
    for col in numeric_columns(df):
        nums = get_numeric_series(df, col)
        if nums.empty:
            continue
        name = str(col).lower()
        if is_raw_pvalue_column(name):
            bad = ((nums < 0) | (nums > 1)).sum()
            if bad:
                add_finding(findings, "P-value bounds", "scientific_bounds", "red", "high", 50, f"Column {col} has {bad} values outside [0, 1].", {"column": col, "bad_count": int(bad)}, column=col)
        if re.search(r"pearson|spearman|correlation|corr", name):
            bad = ((nums < -1) | (nums > 1)).sum()
            if bad:
                add_finding(findings, "Correlation bounds", "scientific_bounds", "red", "high", 50, f"Column {col} has {bad} correlations outside [-1, 1].", {"column": col, "bad_count": int(bad)}, column=col)
        if re.search(r"auc|prob|probability", name):
            bad = ((nums < 0) | (nums > 1)).sum()
            if bad:
                add_finding(findings, "Probability/AUC bounds", "scientific_bounds", "red", "high", 50, f"Column {col} has {bad} values outside [0, 1].", {"column": col, "bad_count": int(bad)}, column=col)
        if role_has(roles, col, "count"):
            bad = (nums < 0).sum()
            if bad:
                add_finding(findings, "Negative counts", "scientific_bounds", "red", "high", 45, f"Column {col} has {bad} negative count-like values.", {"column": col, "bad_count": int(bad)}, column=col)
        inf_count = np.isinf(nums.to_numpy(dtype=float)).sum()
        if inf_count:
            add_finding(findings, "Infinite numeric values", "scientific_bounds", "yellow", "high", 20, f"Column {col} has {inf_count} infinite values.", {"column": col, "inf_count": int(inf_count)}, column=col)


def check_missingness(df: pd.DataFrame, findings: list[Finding]) -> None:
    if df.empty:
        return
    miss = df.isna().mean()
    high_cols = miss[miss >= 0.80]
    if len(high_cols) > 0:
        add_finding(findings, "High-missingness columns", "missingness_structure", "yellow", "medium", 12, f"{len(high_cols)} columns have at least 80% missing values.", {"columns": high_cols.index.astype(str).tolist()[:30]})
    empty_rows = df.isna().all(axis=1).mean()
    if empty_rows >= 0.05:
        add_finding(findings, "All-empty rows", "missingness_structure", "yellow", "medium", 12, f"All-empty rows account for {empty_rows:.1%} of the table.", {"empty_row_fraction": float(empty_rows)})


def audit_table(table: LoadedTable, config: AuditConfig) -> TableAuditResult:
    if table.dataframe is None:
        return TableAuditResult(table.table_name, table.status, table.n_rows, table.n_cols, 0, None, None, {}, [], table.notes)
    df = sample_rows(table.dataframe, config)
    findings: list[Finding] = []
    roles = infer_roles(df)
    missing_fraction = float(df.isna().to_numpy().mean()) if df.size else 0.0
    duplicate_fraction = float(df.duplicated().mean()) if len(df) else 0.0
    check_duplicates(df, findings)
    check_missingness(df, findings)
    check_terminal_digits(df, roles, findings, config)
    check_heaping(df, roles, findings, config)
    check_coordinate_regularities(df, roles, findings, config)
    check_benford(df, roles, findings, config)
    check_scientific_bounds(df, roles, findings)
    return TableAuditResult(
        table_name=table.table_name,
        status=table.status,
        n_rows=table.n_rows,
        n_cols=table.n_cols,
        numeric_cols=len(numeric_columns(df)),
        missing_fraction=missing_fraction,
        duplicate_fraction=duplicate_fraction,
        roles=roles,
        findings=findings,
        notes=table.notes,
    )


def score_file(tables: list[TableAuditResult], config: AuditConfig) -> tuple[float, str]:
    family_scores: dict[str, list[float]] = defaultdict(list)
    for table in tables:
        for finding in table.findings:
            family_scores[finding.family].append(finding.score)
    score = 0.0
    for scores in family_scores.values():
        ordered = sorted(scores, reverse=True)
        score += ordered[0]
        if len(ordered) > 1:
            score += 0.25 * sum(ordered[1:3])
    red_count = sum(1 for t in tables for f in t.findings if f.severity == "red")
    severe_red_count = sum(1 for t in tables for f in t.findings if f.severity == "red" and f.confidence == "high")
    loaded_any = any(t.status == "loaded" for t in tables)
    metadata_only = any(t.status in {"metadata_only", "failed"} for t in tables) and not loaded_any
    if metadata_only:
        return max(score, 25.0), "yellow"
    if score >= config.red_threshold or (red_count >= 3 and score >= 70) or (severe_red_count >= 2 and score >= 70):
        return float(score), "red"
    if score >= config.yellow_threshold:
        return float(score), "yellow"
    return float(score), "green"


def audit_file(path: Path, config: AuditConfig, base_dir: Path) -> FileAuditResult:
    if config.verbose:
        print(f"Auditing {path}")
    tables = load_file(path, config)
    table_results = [audit_table(table, config) for table in tables]
    score, category = score_file(table_results, config)
    status = "loaded" if any(t.status == "loaded" for t in table_results) else table_results[0].status if table_results else "failed"
    notes = [note for t in table_results for note in t.notes]
    rel = str(path.relative_to(base_dir)) if path.is_relative_to(base_dir) else str(path)
    return FileAuditResult(
        path=rel,
        format=file_format(path),
        size_mb=path.stat().st_size / (1024 * 1024),
        status=status,
        category=category,
        score=score,
        tables=table_results,
        notes=notes,
        content_hash=content_hash(path),
    )


def run_cross_file_checks(results: list[FileAuditResult], config: AuditConfig) -> None:
    by_hash: dict[str, list[FileAuditResult]] = defaultdict(list)
    for result in results:
        if result.content_hash:
            by_hash[result.content_hash].append(result)
    for group in by_hash.values():
        if len(group) <= 1:
            continue
        paths = [r.path for r in group]
        for result in group:
            if result.tables:
                finding = Finding(
                    "Identical file content",
                    "cross_file_similarity",
                    "yellow",
                    "medium",
                    18,
                    "This file has identical byte-level content to another audited file.",
                    {"matching_files": paths},
                    ["Identical files can be legitimate when the same input is intentionally reused."],
                )
                result.tables[0].findings.append(finding)
                result.score, result.category = score_file(result.tables, config)


def audit_directory(root: Path, config: AuditConfig, progress_callback: ProgressCallback | None = None) -> list[FileAuditResult]:
    emit_progress(progress_callback, "discovering_files", 0, 0, message=f"Discovering files under {root}")
    files = discover_files(root, config)
    total = len(files)
    if config.verbose:
        print(f"Discovered {total} supported files under {root}")
    emit_progress(progress_callback, "discovered", 0, total, message=f"Discovered {total} supported files.")

    results: list[FileAuditResult] = []
    for index, path in enumerate(files, start=1):
        rel = str(path.relative_to(root)) if path.is_relative_to(root) else str(path)
        emit_progress(progress_callback, "audit_file_start", index - 1, total, rel, f"Auditing {rel}")
        results.append(audit_file(path, config, root))
        emit_progress(progress_callback, "audit_file_done", index, total, rel, f"Finished {rel}")

    emit_progress(progress_callback, "cross_file_checks", total, total, message="Running cross-file checks.")
    run_cross_file_checks(results, config)
    emit_progress(progress_callback, "audit_complete", total, total, message="Audit complete.")
    return results


def generate_simulation_datasets(out_dir: Path) -> list[SimulationExpectedResult]:
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(20260604)
    expected: list[SimulationExpectedResult] = []

    clean = pd.DataFrame(rng.normal(size=(180, 6)), columns=[f"measurement_{i}" for i in range(6)])
    clean["sample_id"] = [f"S{i:03d}" for i in range(len(clean))]
    path = out_dir / "sim_clean_continuous.csv"
    clean.to_csv(path, index=False)
    expected.append(SimulationExpectedResult("clean_continuous", str(path), "green", "not_run", False, "Normal/lognormal-like measurements should be green."))

    days = np.arange(1, 101)
    mono = pd.DataFrame({"day": days, "dose": np.repeat([0, 5, 10, 20], 25), "mouse_weight": 20 + 0.08 * days + rng.normal(0, 0.4, size=100)})
    path = out_dir / "sim_benign_time_dose.csv"
    mono.to_csv(path, index=False)
    expected.append(SimulationExpectedResult("benign_time_dose", str(path), "green", "not_run", False, "Legitimate time/dose structure should not become red."))

    dup = pd.DataFrame(rng.normal(size=(120, 5)), columns=[f"value_{i}" for i in range(5)])
    dup = pd.concat([dup, dup.iloc[20:60], dup.iloc[20:60]], ignore_index=True)
    path = out_dir / "sim_duplicate_blocks.csv"
    dup.to_csv(path, index=False)
    expected.append(SimulationExpectedResult("duplicate_blocks", str(path), "yellow", "not_run", False, "Large duplicated blocks should be at least yellow under conservative red rules."))

    digit = pd.DataFrame({"measurement": np.round(rng.normal(50, 10, 220) * 2) / 2})
    path = out_dir / "sim_terminal_digit_heaping.csv"
    digit.to_csv(path, index=False)
    expected.append(SimulationExpectedResult("terminal_digit_heaping", str(path), "yellow", "not_run", False, "Excess .0/.5 style values should be flagged."))

    rr, cc = np.indices((80, 8))
    grad = pd.DataFrame(10 + 0.3 * rr + 0.7 * cc + rng.normal(0, 0.001, size=(80, 8)), columns=[f"measure_{i}" for i in range(8)])
    path = out_dir / "sim_coordinate_gradient.csv"
    grad.to_csv(path, index=False)
    expected.append(SimulationExpectedResult("coordinate_gradient", str(path), "yellow", "not_run", False, "Coordinate-driven numeric blocks should be at least yellow under conservative red rules."))

    heaped = pd.DataFrame({f"measurement_{i}": rng.choice([1.1, 2.2, 3.3, 4.4, 5.5], size=180) for i in range(4)})
    path = out_dir / "sim_low_randomness_heaped.csv"
    heaped.to_csv(path, index=False)
    expected.append(SimulationExpectedResult("low_randomness_heaped", str(path), "yellow", "not_run", False, "Continuous-looking heaped values should be yellow/red."))

    if openpyxl is not None:
        path = out_dir / "sim_multisheet.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            clean.head(50).to_excel(writer, sheet_name="clean", index=False)
            dup.to_excel(writer, sheet_name="duplicates", index=False)
        expected.append(SimulationExpectedResult("multisheet_excel", str(path), "yellow", "not_run", False, "Excel sheets should be audited when openpyxl is available."))
    return expected


def category_passes(expected: str, observed: str) -> bool:
    if expected == observed:
        return True
    if expected == "yellow" and observed == "red":
        return True
    if expected == "green" and observed == "yellow":
        return True
    return False


def run_simulations(config: AuditConfig, progress_callback: ProgressCallback | None = None) -> tuple[list[FileAuditResult], list[SimulationExpectedResult]]:
    emit_progress(progress_callback, "generating_simulations", 0, 0, message="Generating simulation datasets.")
    expected = generate_simulation_datasets(config.simulation_output_dir)
    sim_config = AuditConfig(**{**asdict(config), "raw_dir": config.simulation_output_dir})
    results = audit_directory(config.simulation_output_dir, sim_config, progress_callback)
    by_path = {str((config.simulation_output_dir / r.path).resolve()): r for r in results}
    for item in expected:
        result = by_path.get(str(Path(item.path).resolve()))
        if result:
            item.observed_category = result.category
            item.passed = category_passes(item.expected_category, result.category)
    return results, expected


def finding_to_dict(finding: Finding) -> dict[str, Any]:
    return asdict(finding)


def result_to_dict(result: FileAuditResult) -> dict[str, Any]:
    data = asdict(result)
    return data


def write_json_report(results: list[FileAuditResult], simulations: list[SimulationExpectedResult], config: AuditConfig, deps: DependencyStatus) -> None:
    if config.output_json is None:
        return
    config.output_json.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "raw_dir": str(config.raw_dir),
        "dependencies": asdict(deps),
        "summary": Counter(r.category for r in results),
        "files": [result_to_dict(r) for r in results],
        "simulations": [asdict(s) for s in simulations],
    }
    with config.output_json.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)


def set_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)


def add_colored_run(paragraph: Any, text: str, category: str) -> None:
    run = paragraph.add_run(text)
    run.bold = True
    if RGBColor is not None:
        colors = {"red": RGBColor(192, 0, 0), "yellow": RGBColor(156, 101, 0), "green": RGBColor(0, 128, 0)}
        run.font.color.rgb = colors.get(category, RGBColor(0, 0, 0))


def add_table_from_rows(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)


def top_findings(result: FileAuditResult, limit: int = 8) -> list[Finding]:
    findings = [f for table in result.tables for f in table.findings]
    return sorted(findings, key=lambda f: f.score, reverse=True)[:limit]


def write_docx_report(results: list[FileAuditResult], simulations: list[SimulationExpectedResult], config: AuditConfig, deps: DependencyStatus) -> None:
    if Document is None:
        raise RuntimeError("python-docx is required to write a Word report. Install with: pip install python-docx")
    config.output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_default_style(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("UniCure Raw Data Audit Report")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph("This report was generated using software from https://github.com/ZexiChen502/Scientific-Raw-Data-Audit.")
    doc.add_paragraph(f"Generated at: {datetime.now().isoformat(timespec='seconds')}")
    doc.add_paragraph(f"Audited directory: {config.raw_dir}")
    doc.add_paragraph("This report is a statistical screening summary. Findings are integrity review indicators requiring contextual interpretation, not proof of misconduct.")

    counts = Counter(r.category for r in results)
    doc.add_heading("Audit function and principles", level=1)
    doc.add_paragraph(
        "This tool performs a statistical screening of raw-data files to identify patterns that may require manual review. "
        "It checks whether tabular data can be read consistently, profiles missingness and duplicates, and evaluates numeric columns for distributional or structural signals that are unusual for measurement-like data."
    )
    principles = [
        "Duplicate and copy/paste indicators: exact or normalized duplicate rows, repeated blocks, duplicate columns, and identical file contents are flagged when they occur at notable levels.",
        "Terminal-digit and rounding patterns: numeric measurements are screened for excessive terminal 0/5 digits, low terminal-digit entropy, and unusually concentrated decimal endings.",
        "Heaping and low-randomness signals: continuous-looking columns are checked for excessive repeated values, low unique-value ratios, and strong rounding patterns.",
        "Table-coordinate regularities: numeric blocks are checked for values that change too regularly with row or column position, including arithmetic progressions, repeated adjacent-cell deltas, and row/column gradient patterns.",
        "Scientific sanity checks: column names are used to test expected bounds, such as p-values in [0, 1], correlations in [-1, 1], non-negative counts, finite numeric values, and severe missingness.",
        "Context-aware false-positive control: columns that look like identifiers, dose/time/rank/epoch variables, coordinates, count data, or statistical summaries are skipped or downgraded where such regularity can be expected.",
    ]
    for item in principles:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Executive summary", level=1)
    add_table_from_rows(doc, ["Category", "Number of files", "Meaning"], [
        ["Red", counts.get("red", 0), "Priority review recommended"],
        ["Yellow", counts.get("yellow", 0), "Contextual review recommended or partial audit"],
        ["Green", counts.get("green", 0), "No major statistical indicators detected"],
    ])

    doc.add_heading("Dependency availability", level=1)
    add_table_from_rows(doc, ["Dependency", "Available", "Purpose"], [
        ["pandas", deps.pandas, "CSV/TSV/TXT/Excel/parquet tables"],
        ["numpy", deps.numpy, "Numeric operations"],
        ["scipy", deps.scipy, "Statistical tests"],
        ["python-docx", deps.python_docx, "Word report"],
        ["openpyxl", deps.openpyxl, "Optional .xlsx support"],
        ["pyreadr", deps.pyreadr, "Optional .Rdata/.rds support"],
        ["anndata", deps.anndata, "Optional .h5ad support"],
        ["h5py", deps.h5py, "Optional .h5 support"],
    ])

    if simulations:
        doc.add_heading("Simulation validation", level=1)
        add_table_from_rows(doc, ["Dataset", "Expected", "Observed", "Pass", "Notes"], [
            [s.name, s.expected_category, s.observed_category, "yes" if s.passed else "no", s.notes]
            for s in simulations
        ])

    doc.add_heading("Top findings", level=1)
    rows = []
    for result in sorted(results, key=lambda r: r.score, reverse=True)[:20]:
        for finding in top_findings(result, 3):
            rows.append([result.category, result.path, finding.check_name, finding.severity, f"{finding.score:.1f}", finding.message])
    add_table_from_rows(doc, ["File category", "File", "Check", "Severity", "Score", "Message"], rows or [["", "", "", "", "", "No findings."]])

    for category, title_text in [("red", "Red files requiring priority review"), ("yellow", "Yellow files requiring contextual review"), ("green", "Green files with no major statistical indicators")]:
        doc.add_heading(title_text, level=1)
        category_results = [r for r in sorted(results, key=lambda x: x.score, reverse=True) if r.category == category]
        if not category_results:
            doc.add_paragraph("None.")
            continue
        for result in category_results:
            p = doc.add_paragraph()
            add_colored_run(p, f"{result.path} — {result.category.upper()} ", result.category)
            p.add_run(f"(score {result.score:.1f}, format {result.format}, size {result.size_mb:.2f} MB)")
            if result.notes:
                doc.add_paragraph("Notes: " + "; ".join(sorted(set(result.notes))[:5]))
            findings = top_findings(result, 6)
            if findings:
                for finding in findings:
                    text = f"{finding.check_name}: {finding.message}"
                    if finding.false_positive_notes:
                        text += " Caveat: " + " ".join(finding.false_positive_notes[:2])
                    doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph("No major findings.", style="List Bullet")

    doc.add_heading("Skipped or metadata-only files", level=1)
    skipped = [r for r in results if r.status in {"metadata_only", "failed"}]
    if skipped:
        add_table_from_rows(doc, ["File", "Status", "Notes"], [[r.path, r.status, "; ".join(sorted(set(r.notes))[:3])] for r in skipped])
    else:
        doc.add_paragraph("None.")

    doc.add_heading("Methods appendix", level=1)
    method_lines = [
        "The audit checks exact/normalized duplicates, repeated blocks, terminal-digit distributions, heaping and low-randomness patterns, coordinate-related regularities, Benford-style first digits where appropriate, scientific bounds, missingness, and cross-file identical content.",
        "The program uses column-role inference to reduce false positives for identifiers, dose/time/rank/epoch columns, coordinates, statistical summaries, and count-like variables.",
        "Red/yellow/green categories combine grouped finding scores across independent check families. A red classification indicates priority review, not proof of fabrication.",
        "Known false-positive contexts include time courses, dose-response designs, sorted rank tables, training curves, count matrices, deliberately rounded measurements, and designed coordinate grids.",
    ]
    for line in method_lines:
        doc.add_paragraph(line)

    doc.save(config.output_docx)


def run_audit(config: AuditConfig, progress_callback: ProgressCallback | None = None) -> tuple[list[FileAuditResult], list[SimulationExpectedResult]]:
    deps = DependencyStatus()
    if not config.no_docx and Document is None:
        raise RuntimeError("python-docx is required unless --no-docx is used. Install with: pip install python-docx")

    all_results: list[FileAuditResult] = []
    simulations: list[SimulationExpectedResult] = []
    if config.run_simulations:
        sim_results, simulations = run_simulations(config, progress_callback)
        all_results.extend(sim_results)
        if not config.audit_simulations_and_raw:
            emit_progress(progress_callback, "writing_reports", len(all_results), len(all_results), message="Writing reports.")
            if not config.no_json:
                write_json_report(all_results, simulations, config, deps)
            if not config.no_docx:
                write_docx_report(all_results, simulations, config, deps)
            emit_progress(progress_callback, "complete", len(all_results), len(all_results), message="Reports written.")
            return all_results, simulations

    raw_results = audit_directory(config.raw_dir, config, progress_callback)
    all_results.extend(raw_results)
    emit_progress(progress_callback, "writing_reports", len(all_results), len(all_results), message="Writing reports.")
    if not config.no_json:
        write_json_report(all_results, simulations, config, deps)
    if not config.no_docx:
        write_docx_report(all_results, simulations, config, deps)
    emit_progress(progress_callback, "complete", len(all_results), len(all_results), message="Reports written.")
    return all_results, simulations


def main() -> None:
    config = parse_args()
    all_results, _ = run_audit(config)
    if config.run_simulations and not config.audit_simulations_and_raw:
        print(f"Audited {len(all_results)} simulation files.")
    else:
        print(f"Audited {len(all_results)} files.")
    if not config.no_docx:
        print(f"Word report: {config.output_docx}")
    if config.output_json is not None and not config.no_json:
        print(f"JSON report: {config.output_json}")


if __name__ == "__main__":
    main()
